"""Unit tests for sprint-planning/generate_sprint_planner.py."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from conftest import sprint_planner as sp


def test_sprint_sheet_layout():
    layout = sp.sprint_sheet_layout(team_count=3, max_backlog_rows=10)
    assert layout["capacity_last"] == sp.CAPACITY_FIRST_DATA_ROW + 2
    assert layout["backlog_last_data"] == layout["backlog_first_data"] + 9


def test_normalize_sprint_length_config():
    days, weeks = sp.normalize_sprint_length_config({"sprint_length": {"weeks": 2}})
    assert weeks == 2.0
    assert days is None

    days2, weeks2 = sp.normalize_sprint_length_config({"sprint_length": {"working_days": 8}})
    assert days2 == 8
    assert weeks2 is None


def test_resolve_sprint_length_working_days():
    cfg = {"sprint_length_weeks": 2.0, "sprint_length_working_days": None}
    assert sp.resolve_sprint_length_working_days(None, cfg) == 10


def test_parse_sprint_numbers_and_next_entries():
    names = ["Sprint 1", "Sprint 2", "Product backlog", "Sprint 10"]
    assert sp.parse_sprint_numbers(names, "Sprint") == [1, 2, 10]

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Sprint 1")
    wb.create_sheet("Sprint 3")
    cfg = {"sprint_name_prefix": "Sprint", "team_members": ["A"], "max_backlog_rows_per_sprint": 5}
    entries = sp.next_sprint_entries(wb, cfg, count=2)
    assert [e["sheet_name"] for e in entries] == ["Sprint 4", "Sprint 5"]


def test_list_formula_and_backlog_formulas():
    assert sp.list_formula(["A", "B"]) == '"A,B"'
    row = 20
    assert "IF" in sp.backlog_completed_sp_formula(row)
    assert "SUM" in sp.committed_sp_sum_formula(10, 15)


def test_update_team_members_in_config(tmp_path: Path):
    cfg_path = tmp_path / "sprint_planner.config.json"
    cfg_path.write_text(
        json.dumps({"team_members": ["Alice", "Bob"], "sprints": 1}),
        encoding="utf-8",
    )
    members = sp.update_team_members_in_config(cfg_path, add="Carol", remove=None)
    assert members == ["Alice", "Bob", "Carol"]


def test_backup_workbook_if_exists_creates_backups_dir(tmp_path: Path):
    workbook = tmp_path / "Sprint-Planning.xlsx"
    workbook.write_bytes(b"placeholder")
    backup = sp.backup_workbook_if_exists(workbook)
    assert backup is not None
    assert backup.parent == tmp_path / sp.BACKUPS_DIR_NAME
    assert backup.is_file()


def _minimal_sprint_config(tmp_path: Path, sprints: int = 2) -> Path:
    cfg_path = tmp_path / "sprint_planner.config.json"
    cfg_path.write_text(
        json.dumps(
            {
                "output_file": "Sprint-Planning.xlsx",
                "team_members": ["Engineer 1", "Engineer 2"],
                "story_point_scale": [1, 2, 3],
                "work_item_types": ["Story"],
                "statuses": ["To Do", "Done"],
                "priorities": ["1 - Must"],
                "sprint_length": {"weeks": 2},
                "max_backlog_rows_per_sprint": 5,
                "product_backlog_rows": 3,
                "sprint_name_prefix": "Sprint",
                "sprints": sprints,
            }
        ),
        encoding="utf-8",
    )
    return cfg_path


def test_generate_addsprint_addengineer(tmp_path: Path):
    cfg_path = _minimal_sprint_config(tmp_path, sprints=1)
    out, _ = sp.generate(cfg_path, force=False)
    assert out.is_file()

    wb = load_workbook(out)
    sprint_name = sp.ordered_sprint_sheet_names(wb)[0]
    ws = wb[sprint_name]
    ws.cell(row=sp.CAPACITY_FIRST_DATA_ROW, column=sp.COL_PTO, value=2)
    ws.cell(
        row=sp.sprint_sheet_layout(2, 5)["backlog_first_data"],
        column=sp.BACKLOG_TITLE,
        value="Test story",
    )
    wb.save(out)

    out2, _ = sp.addsprint(cfg_path, count=1, workbook=None)
    wb2 = load_workbook(out2)
    assert len(sp.ordered_sprint_sheet_names(wb2)) == 2

    out3, backup, _action = sp.mutate_sprint_team(
        cfg_path, add="Engineer 3", remove=None, workbook=None
    )
    assert backup is not None
    assert (tmp_path / sp.BACKUPS_DIR_NAME).is_dir()
    assert "Engineer 3" in json.loads(cfg_path.read_text())["team_members"]

    wb3 = load_workbook(out3)
    ws3 = wb3[sprint_name]
    assert sp.read_capacity_rows(ws3)[-1][0] == "Engineer 3"
    assert sp.read_capacity_rows(ws3)[0][1] == 2

    bf = sp.sprint_sheet_layout(3, 5)["backlog_first_data"]
    assert ws3.cell(row=bf, column=sp.BACKLOG_TITLE).value == "Test story"

    with pytest.raises(ValueError, match="not on the team"):
        sp.mutate_sprint_team(
            cfg_path, add=None, remove="Nobody", workbook=None
        )
