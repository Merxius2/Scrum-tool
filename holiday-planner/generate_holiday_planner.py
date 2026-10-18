#!/usr/bin/env python3
"""Generate team holiday planner Excel workbook from JSON config."""

from __future__ import annotations

import argparse
import colorsys
import json
import random
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_CONFIG_NAME = "holiday_planner.config.json"
DATE_DISPLAY_FORMAT = "%d-%m-%Y"

WEEK_ROW = 1
DAY_ROW = 2
NAME_COL = 1
COUNT_COL = 2
FIRST_DAY_COL = 3
FIRST_DATA_ROW = 3

THIN = Side(style="thin", color="CCCCCC")
WEEK_BORDER = Side(style="medium", color="5B5B5B")
HEADER_FILL = PatternFill("solid", fgColor="E8EEF4")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")


def assign_engineer_colors(count: int, rng: random.Random) -> list[str]:
    """Pick distinct pastel fill colors (RRGGBB) readable with dark text."""
    colors: list[str] = []
    seen: set[str] = set()
    for _ in range(count):
        for _attempt in range(64):
            hue = rng.random()
            saturation = rng.uniform(0.42, 0.78)
            lightness = rng.uniform(0.72, 0.86)
            r, g, b = colorsys.hls_to_rgb(hue, lightness, saturation)
            hex_color = f"{int(r * 255):02X}{int(g * 255):02X}{int(b * 255):02X}"
            if hex_color not in seen:
                seen.add(hex_color)
                colors.append(hex_color)
                break
        else:
            fallback = f"{rng.randint(200, 255):02X}{rng.randint(200, 255):02X}{rng.randint(200, 255):02X}"
            colors.append(fallback)
    return colors


def contrasting_font_color(hex_rgb: str) -> str:
    r = int(hex_rgb[0:2], 16)
    g = int(hex_rgb[2:4], 16)
    b = int(hex_rgb[4:6], 16)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "1F1F1F" if luminance > 0.55 else "FFFFFF"


def cell_border(week_start_columns: set[int], column: int) -> Border:
    left = WEEK_BORDER if column in week_start_columns else THIN
    return Border(left=left, right=THIN, top=THIN, bottom=THIN)


def lock_all_cells(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            cell.protection = Protection(locked=True)


def unlock_cell(ws, row: int, col: int) -> None:
    ws.cell(row=row, column=col).protection = Protection(locked=False)


def enable_sheet_protection(ws) -> None:
    ws.protection.sheet = True
    ws.protection.enable()


def protect_holiday_sheet(
    ws,
    day_columns: list[int],
    last_engineer_row: int,
) -> None:
    """Lock formulas and headers; unlock engineer names and holiday day cells."""
    lock_all_cells(ws)
    for row in range(FIRST_DATA_ROW, last_engineer_row + 1):
        unlock_cell(ws, row, NAME_COL)
        for col in day_columns:
            unlock_cell(ws, row, col)
    enable_sheet_protection(ws)


def backup_workbook_if_exists(path: Path) -> Path | None:
    if not path.is_file():
        return None
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    return backup_path


def engineer_colors_from_workbook(wb, team_count: int) -> list[str] | None:
    """Reuse row label colours from an existing sheet when appending a year."""
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        colors: list[str] = []
        for i in range(team_count):
            row = FIRST_DATA_ROW + i
            cell = ws.cell(row=row, column=NAME_COL)
            fg = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
            if not fg or fg in ("00000000", "0"):
                colors = []
                break
            hex_rgb = fg[2:] if len(fg) == 8 else fg
            colors.append(hex_rgb.upper())
        if len(colors) == team_count:
            return colors
    return None


def resolve_workbook_path(cfg: dict[str, Any], base_dir: Path, workbook: Path | None) -> Path:
    if workbook is None:
        return base_dir / cfg["output_file"]
    if workbook.is_absolute():
        return workbook
    return base_dir / workbook


def iter_weekdays(start: date, end: date):
    d = start
    while d <= end:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def group_by_iso_week(days):
    weeks = []
    current_key = None
    bucket = []
    for d in days:
        key = d.isocalendar()[:2]
        if key != current_key:
            if bucket:
                weeks.append(bucket)
            bucket = [d]
            current_key = key
        else:
            bucket.append(d)
    if bucket:
        weeks.append(bucket)
    return weeks


def parse_date(value: str) -> date:
    return datetime.strptime(value, "%Y-%m-%d").date()


def first_weekday_on_or_after(d: date) -> date:
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def resolve_period(entry: dict[str, Any], today: date) -> tuple[str, date, date]:
    sheet_name = entry.get("sheet_name")
    if not sheet_name:
        raise ValueError('Each year entry needs a "sheet_name".')

    if "year" in entry:
        year = int(entry["year"])
        start = date(year, 1, 1)
        end = date(year, 12, 31)
        return sheet_name, start, end

    start_raw = entry.get("from") or entry.get("start")
    end_raw = entry.get("through") or entry.get("end") or entry.get("to")
    if start_raw is None or end_raw is None:
        raise ValueError(
            f'Period "{sheet_name}" needs either "year" or both "from"/"through" dates.'
        )

    if str(start_raw).lower() == "remainder":
        start = first_weekday_on_or_after(today)
    else:
        start = parse_date(str(start_raw))
        if start.weekday() >= 5:
            start = first_weekday_on_or_after(start)

    end = parse_date(str(end_raw))
    if end < start:
        raise ValueError(f'Period "{sheet_name}": end date is before start date.')

    return sheet_name, start, end


def load_config(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)

    members = data.get("team_members") or data.get("engineers")
    if not members or not isinstance(members, list):
        raise ValueError('"team_members" must be a non-empty list of names.')

    years = data.get("years") or data.get("periods") or data.get("sheets")
    if not years or not isinstance(years, list):
        raise ValueError('"years" must be a non-empty list of period definitions.')

    return {
        "output_file": data.get("output_file", "Team-Holiday-Planner.xlsx"),
        "holiday_marker": data.get("holiday_marker", "H"),
        "color_seed": data.get("color_seed"),
        "team_members": [str(m) for m in members],
        "years": years,
    }


def build_sheet(
    ws,
    title: str,
    start: date,
    end: date,
    team_members: list[str],
    holiday_marker: str,
    engineer_colors: list[str],
):
    ws.title = title[:31]
    days = list(iter_weekdays(start, end))
    if not days:
        raise ValueError(f'Period "{title}" has no weekdays in range.')

    weeks = group_by_iso_week(days)
    last_engineer_row = FIRST_DATA_ROW + len(team_members) - 1
    total_row = last_engineer_row + 1

    ws.cell(row=WEEK_ROW, column=NAME_COL, value="Team member")
    ws.cell(row=DAY_ROW, column=NAME_COL, value="")
    ws.merge_cells(
        start_row=WEEK_ROW,
        start_column=NAME_COL,
        end_row=DAY_ROW,
        end_column=NAME_COL,
    )
    ws.cell(row=WEEK_ROW, column=COUNT_COL, value="Days off")
    ws.cell(row=DAY_ROW, column=COUNT_COL, value="")
    ws.merge_cells(
        start_row=WEEK_ROW,
        start_column=COUNT_COL,
        end_row=DAY_ROW,
        end_column=COUNT_COL,
    )

    col = FIRST_DAY_COL
    day_columns = []
    week_start_columns: set[int] = set()

    for week_days in weeks:
        _, iso_week, _ = week_days[0].isocalendar()
        week_start_col = col
        week_start_columns.add(week_start_col)
        for d in week_days:
            ws.cell(row=DAY_ROW, column=col, value=d.strftime(f"%a {DATE_DISPLAY_FORMAT}"))
            day_columns.append(col)
            col += 1
        week_end_col = col - 1
        week_label = f"W{iso_week:02d} · w/c {week_days[0].strftime(DATE_DISPLAY_FORMAT)}"
        ws.cell(row=WEEK_ROW, column=week_start_col, value=week_label)
        if week_end_col > week_start_col:
            ws.merge_cells(
                start_row=WEEK_ROW,
                start_column=week_start_col,
                end_row=WEEK_ROW,
                end_column=week_end_col,
            )

    last_day_col = col - 1
    first_day_letter = get_column_letter(FIRST_DAY_COL)
    last_day_letter = get_column_letter(last_day_col)
    marker_escaped = holiday_marker.replace('"', '""')

    for i, name in enumerate(team_members):
        row = FIRST_DATA_ROW + i
        color = engineer_colors[i]
        name_cell = ws.cell(row=row, column=NAME_COL, value=name)
        name_cell.fill = PatternFill("solid", fgColor=color)
        name_cell.font = Font(bold=True, color=contrasting_font_color(color))
        ws.cell(
            row=row,
            column=COUNT_COL,
            value=(
                f'=COUNTIF({first_day_letter}{row}:{last_day_letter}{row},'
                f'"{marker_escaped}")'
            ),
        )

    ws.cell(row=total_row, column=NAME_COL, value="Total on holiday")
    ws.cell(row=total_row, column=COUNT_COL, value="")

    for c in day_columns:
        col_letter = get_column_letter(c)
        engineer_range = f"{col_letter}${FIRST_DATA_ROW}:{col_letter}${last_engineer_row}"
        ws.cell(
            row=total_row,
            column=c,
            value=f'=COUNTIF({engineer_range},"{holiday_marker}")',
        )

    ws.cell(
        row=total_row + 1,
        column=NAME_COL,
        value=f'Mark a day off with "{holiday_marker}" in the cell. Names in column A are editable.',
    )
    if last_day_col >= FIRST_DAY_COL:
        ws.merge_cells(
            start_row=total_row + 1,
            start_column=NAME_COL,
            end_row=total_row + 1,
            end_column=min(last_day_col, NAME_COL + 8),
        )

    marker_escaped_dv = holiday_marker.replace('"', '""')
    dv = DataValidation(
        type="list",
        formula1=f'"{marker_escaped_dv}"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = f'Use "{holiday_marker}" for holiday or leave blank.'
    dv.errorTitle = "Holiday marker"
    if day_columns:
        first = get_column_letter(day_columns[0])
        last = get_column_letter(day_columns[-1])
        dv.add(f"{first}{FIRST_DATA_ROW}:{last}{last_engineer_row}")
    ws.add_data_validation(dv)

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions[get_column_letter(NAME_COL)].width = 18
    ws.column_dimensions[get_column_letter(COUNT_COL)].width = 10
    for c in day_columns:
        ws.column_dimensions[get_column_letter(c)].width = 11

    name_header = ws.cell(row=WEEK_ROW, column=NAME_COL)
    name_header.font = header_font
    name_header.alignment = center
    name_header.fill = HEADER_FILL
    name_header.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    count_header = ws.cell(row=WEEK_ROW, column=COUNT_COL)
    count_header.font = header_font
    count_header.alignment = center
    count_header.fill = HEADER_FILL
    count_header.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    for r in (WEEK_ROW, DAY_ROW):
        for c in range(FIRST_DAY_COL, last_day_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = header_font
            cell.alignment = center
            cell.fill = HEADER_FILL
            cell.border = cell_border(week_start_columns, c)

    for r in range(FIRST_DATA_ROW, last_engineer_row + 1):
        ws.cell(row=r, column=NAME_COL).alignment = left
        ws.cell(row=r, column=NAME_COL).border = Border(
            left=THIN, right=THIN, top=THIN, bottom=THIN
        )
        count_cell = ws.cell(row=r, column=COUNT_COL)
        count_cell.alignment = center
        count_cell.font = Font(bold=True)
        count_cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        for c in day_columns:
            cell = ws.cell(row=r, column=c)
            cell.alignment = center
            cell.border = cell_border(week_start_columns, c)

    total_name = ws.cell(row=total_row, column=NAME_COL)
    total_name.font = Font(bold=True)
    total_name.alignment = left
    total_name.fill = TOTAL_FILL
    total_name.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    total_count = ws.cell(row=total_row, column=COUNT_COL)
    total_count.fill = TOTAL_FILL
    total_count.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    for c in day_columns:
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.fill = TOTAL_FILL
        cell.border = cell_border(week_start_columns, c)

    ws.row_dimensions[WEEK_ROW].height = 22
    ws.row_dimensions[DAY_ROW].height = 28
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=FIRST_DAY_COL).coordinate

    if day_columns:
        first = get_column_letter(day_columns[0])
        last = get_column_letter(day_columns[-1])
        count_col_letter = get_column_letter(COUNT_COL)
        count_range = (
            f"{count_col_letter}{FIRST_DATA_ROW}:{count_col_letter}{last_engineer_row}"
        )
        ws.conditional_formatting.add(
            count_range,
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

        marker_cf = holiday_marker.replace('"', '""')
        for i, color in enumerate(engineer_colors):
            row = FIRST_DATA_ROW + i
            row_range = f"{first}{row}:{last}{row}"
            fill = PatternFill("solid", fgColor=color, bgColor=color)
            ws.conditional_formatting.add(
                row_range,
                CellIsRule(
                    operator="equal",
                    formula=[f'"{marker_cf}"'],
                    fill=fill,
                ),
            )

        total_range = f"{first}${total_row}:{last}${total_row}"
        ws.conditional_formatting.add(
            total_range,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="63BE7B",
                mid_type="num",
                mid_value=2,
                mid_color="FFEB84",
                end_type="num",
                end_value=3,
                end_color="F8696B",
            ),
        )

    protect_holiday_sheet(ws, day_columns, last_engineer_row)

    return ws


def normalize_year_entry(entry: Any) -> dict[str, Any]:
    if isinstance(entry, dict):
        return entry
    if isinstance(entry, int):
        return {"sheet_name": str(entry), "year": entry}
    raise ValueError('Each "years" item must be an object or a year number.')


def engineer_colors_for_run(cfg: dict[str, Any], wb=None) -> list[str]:
    team_count = len(cfg["team_members"])
    if wb is not None:
        existing = engineer_colors_from_workbook(wb, team_count)
        if existing:
            return existing
    rng = random.Random(cfg.get("color_seed"))
    return assign_engineer_colors(team_count, rng)


def create_fresh_workbook(cfg: dict[str, Any], today: date) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    engineer_colors = engineer_colors_for_run(cfg)
    for entry in cfg["years"]:
        sheet_name, start, end = resolve_period(normalize_year_entry(entry), today)
        build_sheet(
            wb.create_sheet(),
            sheet_name,
            start,
            end,
            cfg["team_members"],
            cfg["holiday_marker"],
            engineer_colors,
        )
    return wb


def generate(config_path: Path, force: bool = False, today: date | None = None) -> tuple[Path, Path | None]:
    cfg = load_config(config_path)
    today = today or date.today()
    base_dir = config_path.resolve().parent
    out = base_dir / cfg["output_file"]

    if out.is_file() and not force:
        raise ValueError(
            f"{out.name} already exists. Use --force to recreate, or "
            f"'addyear 2028 {out.name}' to append a year sheet."
        )

    backup = backup_workbook_if_exists(out) if out.is_file() else None
    wb = create_fresh_workbook(cfg, today)
    wb.save(out)
    return out, backup


def addyear(
    config_path: Path,
    years: list[int],
    workbook: Path | None,
    today: date | None = None,
) -> tuple[Path, Path | None, int]:
    if not years:
        raise ValueError("addyear requires at least one calendar year (e.g. addyear 2028).")
    cfg = load_config(config_path)
    today = today or date.today()
    base_dir = config_path.resolve().parent
    out = resolve_workbook_path(cfg, base_dir, workbook)

    if not out.is_file():
        raise ValueError(f"Workbook not found: {out}. Run generate first (with or without --force).")

    backup = backup_workbook_if_exists(out)
    wb = load_workbook(out)
    engineer_colors = engineer_colors_for_run(cfg, wb)
    existing = set(wb.sheetnames)
    added = 0

    for year in years:
        entry = {"sheet_name": str(year), "year": year}
        sheet_name = str(year)[:31]
        if sheet_name in existing:
            continue
        sheet_name, start, end = resolve_period(entry, today)
        build_sheet(
            wb.create_sheet(sheet_name),
            sheet_name,
            start,
            end,
            cfg["team_members"],
            cfg["holiday_marker"],
            engineer_colors,
        )
        existing.add(sheet_name)
        added += 1

    if added == 0:
        print("No new year sheets added (sheet names already exist).", file=sys.stderr)
    else:
        wb.save(out)

    return out, backup, added


def main():
    parser = argparse.ArgumentParser(description="Generate team holiday planner Excel file.")
    parser.add_argument(
        "-c",
        "--config",
        type=Path,
        default=None,
        help=f"Path to JSON config (default: {DEFAULT_CONFIG_NAME} next to this script)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Recreate the entire workbook from config (backs up existing file first)",
    )
    parser.add_argument(
        "command",
        nargs="?",
        default="generate",
        choices=("generate", "addyear"),
        help="generate (default): create workbook; addyear: append calendar year sheet(s)",
    )
    parser.add_argument(
        "year_args",
        nargs="*",
        help="With addyear: one or more years (e.g. 2028), optional path to .xlsx last",
    )
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    config_path = args.config or (script_dir / DEFAULT_CONFIG_NAME)
    if not config_path.is_file():
        print(f"Config not found: {config_path}", file=sys.stderr)
        sys.exit(1)

    try:
        backup: Path | None = None
        if args.command == "addyear":
            years: list[int] = []
            workbook: Path | None = None
            for raw in args.year_args:
                if raw.lower().endswith(".xlsx"):
                    workbook = Path(raw)
                    continue
                years.append(int(raw))
            if not years:
                print(
                    "Usage: generate_holiday_planner.py addyear YEAR [YEAR ...] [WORKBOOK.xlsx]",
                    file=sys.stderr,
                )
                sys.exit(1)
            out, backup, added = addyear(config_path, years, workbook)
            if added:
                print(f"Added {added} year sheet(s).", file=sys.stderr)
        else:
            if args.year_args:
                print("Tip: use 'addyear 2028' to append a year to an existing workbook.", file=sys.stderr)
            out, backup = generate(config_path, force=args.force)
    except (ValueError, json.JSONDecodeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        sys.exit(1)

    if backup is not None:
        print(f"Backup: {backup.resolve()}", file=sys.stderr)
    print(f"Saved: {out.resolve()}")


if __name__ == "__main__":
    main()
